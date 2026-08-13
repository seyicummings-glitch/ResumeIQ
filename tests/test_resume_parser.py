import pytest
from app.services.resume_parser import extract_resume_text


def test_extract_text_from_txt():
    file_bytes = b"John Doe\nSoftware Engineer\nPython, FastAPI, PostgreSQL"
    result = extract_resume_text("resume.txt", file_bytes)
    assert "John Doe" in result
    assert "Python" in result


def test_unsupported_file_type_raises_error():
    file_bytes = b"some content"
    with pytest.raises(ValueError):
        extract_resume_text("resume.xyz", file_bytes)


def test_empty_txt_file():
    file_bytes = b""
    result = extract_resume_text("resume.txt", file_bytes)
    assert result == ""


def test_extract_text_from_docx():
    from docx import Document
    import io

    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("Data Scientist, Python, SQL")
    buf = io.BytesIO()
    doc.save(buf)

    result = extract_resume_text("resume.docx", buf.getvalue())
    assert "Jane Doe" in result
    assert "Python" in result


def test_extract_text_from_docx_with_table_layout():
    # Many resume templates put contact info / skills in a table (a sidebar
    # layout, for example) rather than plain paragraphs — python-docx's
    # `.paragraphs` silently skips table content, so this must still surface it.
    from docx import Document
    import io

    doc = Document()
    doc.add_paragraph("Jane Doe")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Phone"
    table.cell(0, 1).text = "+212 6 12 34 56 78"
    table.cell(1, 0).text = "Skills"
    table.cell(1, 1).text = "Python, Docker"
    buf = io.BytesIO()
    doc.save(buf)

    result = extract_resume_text("resume.docx", buf.getvalue())
    assert "+212 6 12 34 56 78" in result
    assert "Python" in result